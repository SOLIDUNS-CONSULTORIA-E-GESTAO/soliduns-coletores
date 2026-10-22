"""Relatório em Excel (identidade SOLIDUNS) e arquivo JSON para a plataforma."""
import datetime as dt
import json

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .util import AVISOS, ORIGEM

ESCURO, OURO, CLARO = "0F181F", "F2C811", "F7F4EA"
F_TIT = Font(name="Calibri", size=16, bold=True, color=ESCURO)
F_SUB = Font(name="Calibri", size=10, italic=True, color="5A6670")
F_CAB = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_TXT = Font(name="Calibri", size=10)
FILL_CAB = PatternFill("solid", fgColor=ESCURO)
FILL_ALT = PatternFill("solid", fgColor=CLARO)
FILL_OURO = PatternFill("solid", fgColor=OURO)
BORDA = Border(bottom=Side(style="thin", color="E2D9B8"))

# coluna -> (título, formato)
COLS = {
    "ticker": ("Ação", None), "empresa": ("Empresa", None), "grupo_nome": ("Setor", None),
    "peso": ("Peso", "0.0%"), "escore": ("Escore S", "0.00"), "preco": ("Preço (R$)", "#,##0.00"),
    "pl": ("P/L", "0.0"), "U_esperado": ("Potencial esperado E[U] (absoluto)", "0.0%"), "U_min": ("Pior caso U_min (absoluto)", "0.0%"),
    "ret_implicito": ("Retorno implícito no preço", "0.0%"), "ke_esperado": ("K_e esperado", "0.0%"),
    "excesso": ("Excesso de retorno", "0.0%"), "excesso_rel": ("Excesso vs. ação mediana", "+0.0%;-0.0%"),
    "U_D": ("Potencial se D", "0.0%"), "U_E": ("Potencial se E", "0.0%"),
    "choque_D": ("Choque imediato se D", "0.0%"), "choque_E": ("Choque imediato se E", "0.0%"),
    "gamma": ("γ (sensib. política)", "0.000"), "g_esperado": ("Cresc. setorial", "0.0%"),
    "igov": ("IGov", "0"), "segmento": ("Segmento B3", None), "estatal": ("Estatal", None),
    "mb_media": ("Margem bruta média", "0.0%"), "q_margem": ("Nota margem Q_MB", "0.00"),
    "z_retorno_esperado": ("z E[U]", "0.00"), "z_pior_caso": ("z U_min", "0.00"),
    "z_crescimento": ("z cresc.", "0.00"), "z_governanca": ("z IGov", "0.00"),
    "beta": ("β", "0.00"), "lambda": ("λ", "0.00"), "sigma": ("Volatilidade", "0.0%"),
    "ke0": ("K_e atual", "0.00%"), "ke_D": ("K_e se D", "0.00%"), "ke_E": ("K_e se E", "0.00%"),
    "g_D": ("g empresa se D", "0.0%"), "g_E": ("g empresa se E", "0.0%"), "payout": ("Payout φ", "0%"),
    "roe": ("ROE", "0.0%"), "pl_justo_D": ("P/L justo se D", "0.0"), "pl_justo_E": ("P/L justo se E", "0.0"),
    "choque_fund_D": ("Choque fundam. D", "0.0%"), "choque_fund_E": ("Choque fundam. E", "0.0%"),
    "choque_emp_D": ("Choque empírico D", "0.0%"), "choque_emp_E": ("Choque empírico E", "0.0%"),
    "gamma_peso_proprio": ("Peso do γ próprio", "0%"), "volume_medio": ("Volume médio (R$)", "#,##0"),
    "valor_mercado": ("Valor de mercado (R$)", "#,##0"),
    "gov_segmento": ("Nota segmento", "0"), "gov_tag_along": ("Nota tag along", "0"),
    "gov_conselho": ("Nota conselho", "0"), "gov_free_float": ("Nota free float", "0"),
    "gov_controle": ("Nota controle", "0"), "gov_historico": ("Nota histórico", "0"),
    "free_float": ("Free float (%)", "0.0"), "pct_independentes": ("Conselho independente", "0%"),
    "ocorrencias": ("Processos CVM (5 anos)", "0"), "controle": ("Controle", None),
    "cnpj": ("CNPJ", None), "motivo": ("Motivo", None), "dt_refer": ("Último balanço", None),
}


def _planilha(wb, nome, titulo, sub, df, colunas):
    ws = wb.create_sheet(nome)
    ws["A1"], ws["A2"] = titulo, sub
    ws["A1"].font, ws["A2"].font = F_TIT, F_SUB
    ws.row_dimensions[1].height = 24
    for j, c in enumerate(colunas, 1):
        cel = ws.cell(row=4, column=j, value=COLS.get(c, (c,))[0])
        cel.font, cel.fill = F_CAB, FILL_CAB
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 32
    for i, r in enumerate(df[colunas].itertuples(index=False), 5):
        for j, (c, v) in enumerate(zip(colunas, r), 1):
            if isinstance(v, (np.floating, float)) and (np.isnan(v) or np.isinf(v)):
                v = None
            if isinstance(v, (np.bool_, bool)):
                v = "sim" if v else "não"
            if isinstance(v, np.generic):
                v = v.item()
            cel = ws.cell(row=i, column=j, value=v)
            cel.font, cel.border = F_TXT, BORDA
            fmt = COLS.get(c, (None, None))[1]
            if fmt:
                cel.number_format = fmt
            if i % 2 == 0:
                cel.fill = FILL_ALT
    for j, c in enumerate(colunas, 1):
        larg = 34 if c in ("empresa", "motivo") else (26 if c in ("grupo_nome", "segmento") else 13)
        ws.column_dimensions[get_column_letter(j)].width = larg
    ws.freeze_panes = "C5"
    return ws


def gerar(caminho, ctx, cfg, ranking, carteiras, nomes_grupo):
    for df in [ranking] + [c[0] for c in carteiras]:
        df["grupo_nome"] = df["grupo"].map(nomes_grupo)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "SOLIDUNS · Robô de Seleção de Carteira"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color=ESCURO)
    ws["A2"] = "Risco Brasil, risco político, crescimento setorial, governança e margem bruta"
    ws["A2"].font = F_SUB
    ws["A3"].fill = FILL_OURO
    linhas = [
        ("Data de referência", ctx["hoje"].strftime("%d/%m/%Y")),
        ("Gerado em", dt.datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Regime do escore", cfg["escore"]["regime"]),
        ("Modo da carteira", cfg["carteira"]["modo"]),
        ("Probabilidade de D (p_D)", ctx["p"]),
        ("Prêmio político Δ (pb)", round(ctx["delta"] * 1e4, 1)),
        ("CDS atual / se D / se E (pb)", f"{ctx['cds0'] * 1e4:.0f} / {ctx['cds_D'] * 1e4:.1f} / {ctx['cds_E'] * 1e4:.1f}"),
        ("K_e da ação média", ctx["ke_m"]),
        ("Selic meta", ctx["selic"] / 100),
        ("Empresas no ranking", len(ranking)),
        ("Empresas excluídas pelos filtros", len(ctx["excluidas"])),
    ]
    r = 5
    for k, v in linhas:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True, size=10)
        c = ws.cell(row=r, column=2, value=v)
        if k.startswith("Probabilidade") or k.startswith("K_e") or k.startswith("Selic"):
            c.number_format = "0.00%"
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Carteiras").font = Font(bold=True, size=12, color=ESCURO)
    r += 1
    cab = ["Carteira", "Ações", "Restrições", "Γ carteira", "Γ da bolsa", "Choque se D", "Choque se E",
           "Excesso vs. mediana", "IGov médio"]
    for j, h in enumerate(cab, 1):
        c = ws.cell(row=r, column=j, value=h)
        c.font, c.fill = F_CAB, FILL_CAB
    for _, res in carteiras:
        r += 1
        vals = [res["rotulo"], res["n"], "ok" if res["restricoes_ok"] else "não atendidas", res["gamma_carteira"],
                res["gamma_bolsa"], res["choque_D"], res["choque_E"], res["excesso_rel"], res["igov_medio"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.number_format = ["@", "0", "@", "0.000", "0.000", "0.0%", "0.0%", "+0.0%;-0.0%", "0"][j - 1]
    r += 2
    ws.cell(row=r, column=1, value="Avisos da execução").font = Font(bold=True, size=12, color=ESCURO)
    for a in AVISOS or ["nenhum"]:
        r += 1
        ws.cell(row=r, column=1, value="• " + a)
    r += 2
    ws.cell(row=r, column=1, value="Material metodológico; não constitui recomendação de investimento.").font = F_SUB
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 16

    c_cart = ["ticker", "empresa", "grupo_nome", "peso", "escore", "preco", "pl", "ret_implicito", "ke_esperado",
              "excesso_rel",
              "choque_D", "choque_E", "gamma", "igov", "segmento", "estatal", "mb_media"]
    for df, res in carteiras:
        _planilha(wb, res["rotulo"][:31], f"Carteira — {res['rotulo']}",
                  f"Modo {res['modo']} · p_D = {res['p_D']:.2f} · Γ = {res['gamma_carteira']:+.3f}", df, c_cart)
    c_rank = ["ticker", "empresa", "grupo_nome", "escore", "z_retorno_esperado", "z_pior_caso", "z_crescimento",
              "z_governanca", "q_margem", "ret_implicito", "ke_esperado", "excesso_rel", "U_esperado", "U_min", "g_esperado", "igov", "mb_media", "pl",
              "volume_medio", "valor_mercado"]
    _planilha(wb, "Ranking", "Ranking completo (fórmula 9.2)", "Escore e fatores padronizados", ranking, c_rank)
    c_cen = ["ticker", "empresa", "grupo_nome", "beta", "lambda", "sigma", "ke0", "ke_D", "ke_E", "roe", "payout",
             "g_D", "g_E", "pl", "pl_justo_D", "pl_justo_E", "U_D", "U_E", "gamma", "gamma_peso_proprio",
             "choque_fund_D", "choque_fund_E", "choque_emp_D", "choque_emp_E", "choque_D", "choque_E", "dt_refer"]
    _planilha(wb, "Cenários por ação", "Custo de capital, valor e choques por cenário",
              "Fórmulas 3.3, 3.6, 4.6, 4.7, 8.1 a 8.7, 10.2 e 10.3", ranking, c_cen)
    c_gov = ["ticker", "empresa", "igov", "segmento", "controle", "free_float", "pct_independentes", "ocorrencias",
             "gov_segmento", "gov_tag_along", "gov_conselho", "gov_free_float", "gov_controle", "gov_historico"]
    _planilha(wb, "Governança", "Índice SOLIDUNS de Governança (fórmula 6.1)", "Notas de 0 a 100 por componente",
              ranking, c_gov)
    par = pd.DataFrame([(k, str(v[0]), v[1]) for k, v in ORIGEM.items()], columns=["parametro", "valor", "origem"])
    ws2 = wb.create_sheet("Parâmetros")
    ws2["A1"], ws2["A1"].font = "Parâmetros usados nesta execução", F_TIT
    for j, h in enumerate(["Parâmetro", "Valor", "Origem"], 1):
        c = ws2.cell(row=3, column=j, value=h)
        c.font, c.fill = F_CAB, FILL_CAB
    for i, row in enumerate(par.itertuples(index=False), 4):
        for j, v in enumerate(row, 1):
            ws2.cell(row=i, column=j, value=v).font = F_TXT
    ws2.column_dimensions["A"].width, ws2.column_dimensions["B"].width, ws2.column_dimensions["C"].width = 36, 18, 90
    pq = ctx["pesquisas"][["data_divulgacao", "instituto", "amostra", "pct_D", "pct_E", "base", "v_D"]]
    ws3 = wb.create_sheet("Pesquisas")
    ws3["A1"], ws3["A1"].font = "Pesquisas usadas (pesquisas.csv)", F_TIT
    for j, h in enumerate(["Divulgação", "Instituto", "Amostra", "% D", "% E", "Base", "v_D"], 1):
        c = ws3.cell(row=3, column=j, value=h)
        c.font, c.fill = F_CAB, FILL_CAB
    for i, row in enumerate(pq.itertuples(index=False), 4):
        for j, v in enumerate(row, 1):
            ws3.cell(row=i, column=j, value=v.item() if isinstance(v, np.generic) else v)
    ws3.column_dimensions["B"].width = 26
    _planilha(wb, "Excluídas", "Empresas fora do universo", "Filtros da seção 10.1", ctx["excluidas"],
              ["cnpj", "empresa", "motivo"])
    wb.save(caminho)


def gerar_json(caminho, ctx, carteiras, ranking):
    def limpo(df, cols):
        return json.loads(df[cols].replace({np.nan: None}).to_json(orient="records", force_ascii=False))
    saida = {
        "data_referencia": ctx["hoje"].isoformat(), "p_D": ctx["p"], "delta_pb": ctx["delta"] * 1e4,
        "cds_pb": {"atual": ctx["cds0"] * 1e4, "D": ctx["cds_D"] * 1e4, "E": ctx["cds_E"] * 1e4},
        "carteiras": [{"resumo": res, "acoes": limpo(df, ["ticker", "empresa", "peso", "escore", "choque_D",
                                                          "choque_E", "gamma"])} for df, res in carteiras],
        "ranking": limpo(ranking, ["ticker", "empresa", "grupo", "escore", "U_esperado", "U_min", "g_esperado",
                                   "igov", "mb_media", "gamma"]),
        "avisos": AVISOS,
    }
    json.dump(saida, open(caminho, "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)
