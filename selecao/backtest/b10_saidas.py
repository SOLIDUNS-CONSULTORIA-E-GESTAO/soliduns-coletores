"""Gráficos (PNG) e planilha Excel do backtest."""
import sys, json
import numpy as np, pandas as pd
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
B, V, V2, OUT = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]   # V = configuração atual, V2 = configuração original
nav = pd.read_pickle(f'{V}/bt_nav.pkl'); cart = pd.read_pickle(f'{V}/bt_carteiras.pkl'); a = pd.read_pickle(f'{V}/bt_analise.pkl')
nav2 = pd.read_pickle(f'{V2}/bt_nav.pkl'); cart2 = pd.read_pickle(f'{V2}/bt_carteiras.pkl'); a2 = pd.read_pickle(f'{V2}/bt_analise.pkl')
C0, I, CDI, EU, EE = list(nav.columns)
C, CO = 'Carteira SOLIDUNS – config. atual', 'Carteira SOLIDUNS – config. original'
nav = nav.rename(columns={C0: C}); nav.insert(1, CO, nav2[C0])
for k in ('ciclos', 'regimes_selic', 'eleicoes', 'risco_pais'):
    a[k] = a[k].rename(columns={C0: C}); a[k].insert(list(a[k].columns).index(C) + 1, CO, a2[k][C0].values)
m2 = a2['metricas']; m2 = m2[m2.serie == C0].assign(serie=CO); a['metricas'] = pd.concat([a['metricas'].replace({C0: C}), m2]).sort_values(['periodo'], kind='stable')
COR = {C: '#2a78d6', CO: '#e87ba4', I: '#eb6834', CDI: '#1baf7a', EE: '#eda100', EU: '#6250d6'}
CURTO = {C: 'Carteira (config. atual)', CO: 'Carteira (config. original)', I: 'Ibovespa', CDI: 'CDI', EE: 'Elegíveis (peso igual)', EU: 'Universo líquido (peso igual)'}
INK, INK2, GRID = '#0b0b0b', '#52514e', '#e6e5e1'
plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 9, 'axes.edgecolor': GRID, 'axes.labelcolor': INK2,
                     'xtick.color': INK2, 'ytick.color': INK2, 'axes.spines.top': False, 'axes.spines.right': False,
                     'axes.grid': True, 'grid.color': GRID, 'grid.linewidth': 0.6, 'axes.axisbelow': True, 'figure.dpi': 200})
pct = FuncFormatter(lambda v, _: f'{v:.0%}')

# 1. patrimônio (escala log), com faixas de queda da Selic
fig, ax = plt.subplots(figsize=(7.2, 3.6))
for c in a['ciclos'].itertuples():
    if c.tipo == 'Queda': ax.axvspan(pd.Timestamp(c.inicio), pd.Timestamp(c.fim), color='#cde2fb', alpha=0.55, lw=0)
for s in (EE, CDI, I, CO, C):
    v = nav[s] / nav[s].iloc[0] * 100
    ax.plot(v.index, v, color=COR[s], lw=2.2 if s == C else 1.5)
    ax.annotate(f'{CURTO[s]}  {v.iloc[-1]:,.0f}'.replace(',', '.'), (v.index[-1], v.iloc[-1]), xytext=(4, 0), textcoords='offset points',
                va='center', fontsize=7.5, color=INK)
ax.set_yscale('log'); ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f'{v:,.0f}'.replace(',', '.')))
ax.set_xlim(nav.index[0], nav.index[-1] + pd.Timedelta(days=2600)); ax.set_xticks([pd.Timestamp(f'{y}-01-01') for y in range(2002, 2027, 4)], [str(y) for y in range(2002, 2027, 4)]); ax.set_ylabel('R$ 100 aplicados em jan/2001 (escala log)')
ax.text(0.01, 0.97, 'Faixas azuis: ciclos de queda da Selic', transform=ax.transAxes, fontsize=7.5, color=INK2, va='top')
fig.tight_layout(); fig.savefig(f'{OUT}/g1_patrimonio.png'); plt.close(fig)

# 2. queda máxima acumulada (drawdown)
fig, ax = plt.subplots(figsize=(7.2, 2.6))
for s in (I, CO, C):
    dd = nav[s] / nav[s].cummax() - 1; ax.plot(dd.index, dd, color=COR[s], lw=1.5, label=CURTO[s])
ax.yaxis.set_major_formatter(pct); ax.legend(frameon=False, loc='lower left', fontsize=8); ax.set_ylabel('Distância do pico')
fig.tight_layout(); fig.savefig(f'{OUT}/g2_drawdown.png'); plt.close(fig)

# 3. ciclos de queda: retorno acumulado no ciclo
q = a['ciclos'][a['ciclos'].tipo == 'Queda'].reset_index(drop=True)
def _n(t): return f"{float(t.strip().rstrip('%')):g}".replace('.', ',')
lab = [f"{pd.Timestamp(r.inicio):%m/%y}–{pd.Timestamp(r.fim):%m/%y}\n" + '→'.join(_n(t) for t in r.selic.split('→')) for r in q.itertuples()]
fig, ax = plt.subplots(figsize=(7.2, 3.2)); x = np.arange(len(q)); w = 0.2
for k, s in enumerate((C, CO, I, CDI)):
    ax.bar(x + (k - 1.5) * w, q[s], w * 0.92, color=COR[s], label=CURTO[s])
ax.set_xticks(x, lab, fontsize=6.2); ax.set_xlabel('Período do ciclo e Selic meta no início → no fim (% a.a.)', fontsize=7.5); ax.yaxis.set_major_formatter(pct); ax.axhline(0, color=INK2, lw=0.6)
ax.legend(frameon=False, ncol=2, fontsize=7.5, loc='upper right'); ax.set_ylabel('Retorno no ciclo')
fig.tight_layout(); fig.savefig(f'{OUT}/g3_ciclos_queda.png'); plt.close(fig)

# 4. retorno ao ano por regime da Selic
rg = a['regimes_selic'].set_index('regime')
fig, ax = plt.subplots(figsize=(7.2, 2.8)); x = np.arange(3); w = 0.2
for k, s in enumerate((C, CO, I, CDI)):
    b = ax.bar(x + (k - 1.5) * w, rg[s], w * 0.92, color=COR[s], label=CURTO[s])
    for xi, v in zip(x + (k - 1.5) * w, rg[s]): ax.text(xi, v + (0.008 if v >= 0 else -0.03), f'{v:.0%}', ha='center', fontsize=6.8, color=INK)
ax.set_xticks(x, ['Selic em queda', 'Selic em alta', 'Selic estável']); ax.yaxis.set_major_formatter(pct); ax.axhline(0, color=INK2, lw=0.6)
ax.legend(frameon=False, ncol=4, fontsize=7.5, loc='upper right'); ax.set_ylabel('Retorno ao ano')
fig.tight_layout(); fig.savefig(f'{OUT}/g4_regimes.png'); plt.close(fig)

json.dump({'ok': True}, open(f'{OUT}/g.json', 'w'))

# Excel
with pd.ExcelWriter(f'{OUT}/Backtest_Carteira_SOLIDUNS_2001_2026.xlsx', engine='openpyxl') as xw:
    m = a['metricas'].copy(); m.columns = ['Período', 'Série', 'Retorno ao ano', 'Volatilidade ao ano', 'Sharpe (vs CDI)', 'Queda máxima', 'Retorno acumulado']
    m.to_excel(xw, sheet_name='Métricas', index=False)
    a['ciclos'].rename(columns={'tipo': 'Ciclo da Selic', 'inicio': 'Início', 'fim': 'Fim', 'selic': 'Selic meta', 'meses': 'Meses'}).to_excel(xw, sheet_name='Ciclos da Selic', index=False)
    a['regimes_selic'].rename(columns={'regime': 'Regime', 'pregoes': 'Pregões'}).to_excel(xw, sheet_name='Regimes (ao ano)', index=False)
    a['risco_pais'].rename(columns={'regime': 'Trimestres'}).to_excel(xw, sheet_name='Risco-país (média trim.)', index=False)
    a['eleicoes'].rename(columns={'ano': 'Ano (jul–dez)'}).to_excel(xw, sheet_name='Eleições', index=False)
    a['setores_regime'].reset_index().rename(columns={'grupo': 'Grupo setorial'}).to_excel(xw, sheet_name='Setores por regime (atual)', index=False)
    a2['setores_regime'].reset_index().rename(columns={'grupo': 'Grupo setorial'}).to_excel(xw, sheet_name='Setores por regime (orig.)', index=False)
    pd.concat([a['custos'].assign(config='atual'), a2['custos'].assign(config='original')]).rename(columns={'custo_por_lado': 'Custo por giro', 'cagr': 'Retorno ao ano'}).to_excel(xw, sheet_name='Custos', index=False)
    pd.DataFrame([{'Configuração': cf, 'Teste': k, **{kk: float(vv) for kk, vv in v.items()}} for cf, aa in (('atual', a), ('original', a2)) for k, v in (('Total', aa['aleatorias']), ('Fase B', aa['aleatorias_B']), ('Fase A', aa['aleatorias_A']))]).to_excel(xw, sheet_name='Carteiras aleatórias', index=False)
    an = nav.resample('YE').last(); an = an.pct_change().fillna(an.iloc[0] / nav.iloc[0] - 1); an.index = an.index.year
    an.reset_index().rename(columns={'data': 'Ano', 'index': 'Ano'}).to_excel(xw, sheet_name='Retorno por ano', index=False)
    h = cart[['data', 'fase', 'ticker', 'setor', 'grupo', 'estatal', 'peso', 'escore', 'igov', 'pl', 'roe', 'mb', 'ke']].copy()
    h['data'] = h['data'].dt.date
    h.columns = ['Data', 'Fase', 'Ação', 'Setor CVM', 'Grupo', 'Estatal', 'Peso', 'Escore', 'IGov', 'P/L', 'ROE', 'Margem bruta', 'K_e']
    h.to_excel(xw, sheet_name='Carteiras (config. atual)', index=False)
    h2 = cart2[['data', 'fase', 'ticker', 'setor', 'grupo', 'estatal', 'peso', 'escore', 'igov', 'pl', 'roe', 'mb', 'ke']].copy(); h2['data'] = h2['data'].dt.date
    h2.columns = h.columns; h2.to_excel(xw, sheet_name='Carteiras (config. original)', index=False)
    nv = nav.copy(); nv.index = nv.index.date; nv.index.name = 'Data'; nv.reset_index().to_excel(xw, sheet_name='Patrimônio diário', index=False)
    for ws in xw.book.worksheets:
        ws.freeze_panes = 'A2'
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(10, min(38, max(len(str(c.value or '')) for c in col[:60]) + 2))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                h0 = str(ws.cell(1, c.column).value)
                if isinstance(c.value, float):
                    c.number_format = '0.00' if h0 in ('Escore', 'IGov', 'P/L', 'Sharpe (vs CDI)', 'Retorno acumulado') else '0.0%'
                    if ws.title == 'Patrimônio diário': c.number_format = '#,##0.00'
        from openpyxl.styles import Font, PatternFill
        for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='0F181F')
print('ok')
