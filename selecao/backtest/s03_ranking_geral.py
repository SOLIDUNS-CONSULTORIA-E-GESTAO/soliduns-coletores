"""2º ranking: todas as empresas elegíveis, da melhor para a pior, numa lista única.
Nota final = média de (nota dentro do setor) e (nota contra todo o mercado do mesmo tipo: financeiras x não financeiras).
Setores com menos de 3 empresas usam só a nota de mercado. Alertas de endividamento e lucro aparecem, mas não entram na nota."""
import sys
import numpy as np, pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
B, O = sys.argv[1], sys.argv[2]
FORA_GRUPOS = sys.argv[3].split(',') if len(sys.argv) > 3 else []
E = pd.read_pickle(f'{B}/setorial/carteira_setorial.pkl')
E = E[~E['grupo'].isin(FORA_GRUPOS)].copy()
RJ = E[E['empresa'].str.contains('RECUPERA', case=False, na=False)]
E = E[~E.index.isin(RJ.index)]
X = pd.read_pickle(f'{B}/setorial/carteira_setorial_fora.pkl')
X = X[~X['grupo'].isin(FORA_GRUPOS)]
X = pd.concat([X, RJ[['ticker', 'empresa', 'grupo', 'setor_cvm', 'lucro_liq_ttm', 'patrimonio_liq', 'receita_ttm']].assign(so_holding=False, rj=True)])
NOMES = {'agro_exportador': 'Agro exportador', 'bancos': 'Bancos e financeiras', 'consumo_basico': 'Consumo básico', 'energia_saneamento': 'Energia e saneamento',
         'holdings': 'Holdings', 'industria': 'Indústria (bens de capital)', 'lazer_turismo': 'Lazer e turismo', 'mineracao_siderurgia': 'Mineração e siderurgia',
         'papel_celulose': 'Papel e celulose', 'petroleo': 'Petróleo e gás', 'petroquimica': 'Petroquímica', 'saude_educacao': 'Saúde e educação',
         'seguradoras': 'Seguradoras', 'tecnologia_telecom': 'Tecnologia e telecom', 'transporte': 'Transporte e logística'}
IND = {'nao_financeira': ['liq_corrente', 'margem_bruta', 'giro_ativo', 'renovacao', 'cagr_receita'],
       'financeira': ['pl_ativo', 'margem_liq', 'roe', 'cagr_receita', 'cagr_lucro']}
E['nota_mercado'] = np.nan
for t, cols in IND.items():
    m = E['tipo'] == t
    E.loc[m, 'nota_mercado'] = (E.loc[m, cols].rank(pct=True).mean(axis=1) * 100).round(1)
n_setor = E.groupby('grupo')['ticker'].transform('size')
E['nota_setor'] = np.where(n_setor >= 3, E['nota'], np.nan)
E['nota_final'] = E[['nota_setor', 'nota_mercado']].mean(axis=1).round(1)
E = E.sort_values(['nota_final', 'nota_mercado'], ascending=False).reset_index(drop=True)
E['Nº'] = np.arange(1, len(E) + 1)

# alertas (não entram na nota)
ebitda = E['ebit_ttm'] + E['da'].fillna(0)
E['dl_ebitda'] = np.where((E['tipo'] == 'nao_financeira') & (ebitda > 0), (E['divida_bruta'] - E['caixa']) / ebitda, np.nan)
E['payout'] = E['dividendos_ttm'] / E['lucro_liq_ttm']
def alertas(r):
    a = []
    if r.tipo == 'nao_financeira' and pd.notna(r.dl_ebitda) and r.dl_ebitda > 3: a.append('dívida líquida > 3x EBITDA')
    if pd.notna(r.cagr_lucro) and r.cagr_lucro < 0: a.append('lucro em queda (5 anos)')
    if r.tipo == 'nao_financeira' and pd.notna(r.liq_corrente) and r.liq_corrente < 1: a.append('liquidez corrente < 1')
    if pd.notna(r.payout) and r.payout > 1: a.append('dividendos acima do lucro')
    if r.ticker in ('ITSA4', 'ITUB4'): a.append('Itaúsa controla o Itaú: mesma exposição')
    return '; '.join(a)
E['Alertas'] = E.apply(alertas, axis=1)
E['Setor'] = E['grupo'].map(NOMES)

cols = {'Nº': 'Nº', 'ticker': 'Ação', 'empresa': 'Empresa', 'Setor': 'Setor', 'nota_final': 'Nota final (0–100)', 'nota_setor': 'Nota no setor',
        'nota_mercado': 'Nota no mercado', 'liq_corrente': 'Liquidez corrente', 'margem_bruta': 'Margem bruta', 'giro_ativo': 'Giro do ativo',
        'renovacao': 'Renovação (capex/deprec.)', 'cagr_receita': 'CAGR receita 5a', 'pl_ativo': 'PL/Ativo', 'margem_liq': 'Margem líquida', 'roe': 'ROE',
        'cagr_lucro': 'CAGR lucro 5a', 'dl_ebitda': 'Dív. líquida/EBITDA', 'Alertas': 'Alertas (fora da nota)', 'dt_refer': 'Último balanço'}
R = E.rename(columns=cols)[list(cols.values())]
PCT = {'Margem bruta', 'CAGR receita 5a', 'PL/Ativo', 'Margem líquida', 'ROE', 'CAGR lucro 5a'}
X = X.assign(Motivo=np.where(X.get('rj', False) == True, 'em recuperação judicial', np.where(X['so_holding'], 'holding sem receita operacional',
                             np.where(X['patrimonio_liq'] <= 0, 'patrimônio líquido negativo', 'prejuízo nos últimos 12 meses'))),
             Setor=X['grupo'].map(NOMES))[['ticker', 'empresa', 'Setor', 'Motivo', 'lucro_liq_ttm']]
X.columns = ['Ação', 'Empresa', 'Setor', 'Motivo', 'Lucro 12 meses (R$)']
VERDE, VERDE_F = PatternFill('solid', fgColor='C6EFCE'), Font(color='006100', bold=True)
arq = f'{O}/Ranking_Geral_Acoes_SOLIDUNS.xlsx'
with pd.ExcelWriter(arq, engine='openpyxl') as xw:
    R.to_excel(xw, sheet_name='Ranking geral', index=False)
    X.to_excel(xw, sheet_name='Fora do ranking', index=False)
    pd.DataFrame({'Item': ['Ordem', 'Nota no setor', 'Nota no mercado', 'Indicadores (não financeiras)', 'Indicadores (bancos e seguradoras)',
                           'Destaque', 'Alertas', 'Universo', 'Setores fora', 'Dados'],
                  'Regra': ['Da melhor para a pior pela nota final = média da nota no setor e da nota no mercado',
                            'Média dos percentis dos indicadores entre as empresas do mesmo setor (só para setores com 3 ou mais empresas)',
                            'Média dos percentis dos indicadores entre todas as empresas do mesmo tipo (não financeiras ou financeiras)',
                            'Liquidez corrente, margem bruta, giro do ativo, renovação de ativos (capex/depreciação 2025) e CAGR de receita 5 anos',
                            'Patrimônio/Ativo, margem líquida, ROE, CAGR de receita e CAGR de lucro 5 anos',
                            'As 20 primeiras em verde',
                            'Dívida líquida/EBITDA acima de 3, lucro em queda, liquidez corrente abaixo de 1 e dividendos acima do lucro: informados, sem entrar na nota',
                            'Ações com volume médio ≥ R$ 2 milhões/dia (63 pregões até 21/10/2026), uma classe por empresa, lucro e patrimônio positivos',
                            'Varejo, construção civil, saúde e educação, transporte e logística',
                            'CVM: ITR e DFP entregues até 22/10/2026']}).to_excel(xw, sheet_name='Metodologia', index=False)
    for ws in xw.book.worksheets:
        ws.freeze_panes = 'C2' if ws.title == 'Ranking geral' else 'A2'
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='0F181F')
            c.alignment = Alignment(wrap_text=True, vertical='center')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(9, min(48, max(len(str(c.value or '')) for c in col[1:]) + 2 if len(col) > 1 else 12))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                h = str(ws.cell(1, c.column).value)
                if isinstance(c.value, float):
                    c.number_format = '0.0%' if h in PCT else ('#,##0' if 'R$' in h else '0.0' if 'Nota' in h else '0.00')
        if ws.title == 'Ranking geral':
            for row in ws.iter_rows(min_row=2, max_row=21):
                for c in row: c.fill = VERDE; c.font = VERDE_F
print(len(R), 'empresas;', len(X), 'fora')
print(R[['Nº', 'Ação', 'Empresa', 'Setor', 'Nota final (0–100)', 'Alertas (fora da nota)']].to_string(index=False))
