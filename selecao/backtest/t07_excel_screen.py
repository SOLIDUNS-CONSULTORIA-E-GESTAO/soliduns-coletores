"""Planilha da varredura de alvo 2 em toda a B3."""
import sys
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
T, OUT = sys.argv[1], sys.argv[2]
R = pd.read_pickle(f'{T}/screen_b3.pkl').copy()
NOMES = {'agro_exportador': 'Agro exportador', 'bancos': 'Bancos e financeiras', 'consumo_basico': 'Consumo básico', 'energia_saneamento': 'Energia e saneamento',
         'holdings': 'Holdings', 'industria': 'Indústria', 'lazer_turismo': 'Lazer e turismo', 'mineracao_siderurgia': 'Mineração e siderurgia', 'papel_celulose': 'Papel e celulose',
         'petroleo': 'Petróleo e gás', 'petroquimica': 'Petroquímica', 'saude_educacao': 'Saúde e educação', 'seguradoras': 'Seguradoras', 'tecnologia_telecom': 'Tecnologia e telecom',
         'transporte': 'Transporte e logística', 'varejo': 'Varejo', 'construcao': 'Construção civil', 'sem_cadastro': '—', 'nao_mapeado': '—'}
R['Setor'] = R.grupo_orig.map(NOMES)
R['ganho_atual'] = R.alvo2 / R.preco - 1
R.insert(0, 'Nº', range(1, len(R) + 1))
cols = {'Nº': 'Nº', 'empresa': 'Empresa', 'Setor': 'Setor', 'preco': 'Preço 21/10', 'ent_ideal': 'Entrada ideal', 'ent_max': 'Entrada máxima', 'stop': 'Stop',
        'alvo1': 'Alvo 1', 'alvo2': 'Alvo 2', 'ganho2': 'Ganho até alvo 2 (da entrada ideal)', 'ganho_atual': 'Ganho até alvo 2 (do preço atual)', 'perda': 'Perda até stop',
        'rr2': 'Ganho / risco', 'ret_0510': 'Alta em 05/10', 'ret21': 'Alta em 21 dias', 'ifr': 'IFR 14', 'adx': 'ADX', 'alertas': 'Alertas'}
X = R.reset_index().rename(columns={'index': 'Ação'})
X = X[['Ação'] + list(cols)].rename(columns=cols)
X = X[['Nº', 'Ação'] + [c for c in X.columns if c not in ('Nº', 'Ação')]]
PCT = {'Ganho até alvo 2 (da entrada ideal)', 'Ganho até alvo 2 (do preço atual)', 'Perda até stop', 'Alta em 05/10', 'Alta em 21 dias'}
VERDE, VERDE_F, CINZA_F = PatternFill('solid', fgColor='C6EFCE'), Font(color='006100', bold=True), Font(color='8A8F94')
with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
    X.to_excel(xw, sheet_name='Alvo 2 – toda a B3', index=False)
    ws = xw.book['Alvo 2 – toda a B3']
    ws.freeze_panes = 'C2'
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='0F181F'); c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 42
    larg = {'Nº': 5, 'Ação': 9, 'Empresa': 40, 'Setor': 20, 'Alertas': 70}
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = larg.get(col[0].value, 12)
    for i, row in enumerate(ws.iter_rows(min_row=2), start=0):
        r = R.iloc[i]
        for c in row:
            h = ws.cell(1, c.column).value
            if isinstance(c.value, float): c.number_format = '0.0%' if h in PCT else ('0.0' if h in ('IFR 14', 'ADX', 'Ganho / risco') else '#,##0.00')
            if r.ganho2 >= 0.20 and r.limpa: c.fill = VERDE; c.font = VERDE_F
            elif not r.limpa: c.font = CINZA_F
print(len(X), (R.ganho2 >= .2).sum(), ((R.ganho2 >= .2) & R.limpa).sum())
