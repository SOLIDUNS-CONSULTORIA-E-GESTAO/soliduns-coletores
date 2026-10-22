"""Planilha da carteira teórica setorial (4 melhores por setor) a partir de carteira_setorial.pkl."""
import sys
import pandas as pd
from openpyxl.styles import Font, PatternFill
B, O = sys.argv[1], sys.argv[2]
FORA = sys.argv[3].split(',') if len(sys.argv) > 3 else []   # grupos retirados da lista por decisão de gestão
E = pd.read_pickle(f'{B}/setorial/carteira_setorial.pkl')
E = E[~E['grupo'].isin(FORA)]
NOMES = {'agro_exportador': 'Agro exportador', 'bancos': 'Bancos e financeiras', 'consumo_basico': 'Consumo básico', 'energia_saneamento': 'Energia e saneamento',
         'holdings': 'Holdings', 'industria': 'Indústria (bens de capital)', 'lazer_turismo': 'Lazer e turismo', 'mineracao_siderurgia': 'Mineração e siderurgia',
         'papel_celulose': 'Papel e celulose', 'petroleo': 'Petróleo e gás', 'petroquimica': 'Petroquímica', 'saude_educacao': 'Saúde e educação',
         'seguradoras': 'Seguradoras', 'tecnologia_telecom': 'Tecnologia e telecom', 'transporte': 'Transporte e logística'}
E['Setor'] = E.grupo.map(NOMES)
T = E[E.pos <= 4].copy(); T['Peso teórico'] = 1 / len(T)
cols = {'Setor': 'Setor', 'pos': 'Posição', 'ticker': 'Ação', 'empresa': 'Empresa', 'nota': 'Nota (0–100)', 'liq_corrente': 'Liquidez corrente',
        'margem_bruta': 'Margem bruta', 'giro_ativo': 'Giro do ativo', 'renovacao': 'Renovação (capex/deprec.)', 'cagr_receita': 'CAGR receita 5a',
        'pl_ativo': 'PL/Ativo', 'margem_liq': 'Margem líquida', 'roe': 'ROE', 'cagr_lucro': 'CAGR lucro 5a', 'setor_cvm': 'Setor CVM', 'dt_refer': 'Último balanço'}
A = E.rename(columns=cols)[list(cols.values())]
PCT = ('Margem bruta', 'CAGR receita 5a', 'PL/Ativo', 'Margem líquida', 'ROE', 'CAGR lucro 5a', 'Peso teórico')
with pd.ExcelWriter(f'{O}/Carteira_Setorial_Fundamentos_SOLIDUNS.xlsx', engine='openpyxl') as xw:
    T.rename(columns=cols)[list(cols.values()) + ['Peso teórico']].sort_values(['Setor', 'Posição']).to_excel(xw, sheet_name='Carteira (4 por setor)', index=False)
    A.sort_values(['Setor', 'Posição']).to_excel(xw, sheet_name='Ranking completo', index=False)
    pd.DataFrame({'Item': ['Universo', 'Filtros', 'Não financeiras', 'Financeiras (bancos e seguradoras)', 'Nota', 'Renovação', 'Dados', 'Excluídos por decisão', 'Observação'],
                  'Regra': ['Ações com volume médio ≥ R$ 2 milhões/dia nos últimos 63 pregões (até 21/10/2026), uma classe por empresa',
                            'Lucro líquido dos últimos 12 meses positivo e patrimônio líquido positivo; holdings sem receita operacional fora',
                            'Liquidez corrente, margem bruta, giro do ativo (receita/ativo total), renovação de ativos e CAGR de receita de 5 anos',
                            'Patrimônio/Ativo (no lugar da liquidez corrente), margem líquida, ROE, CAGR de receita e CAGR de lucro de 5 anos',
                            'Média dos percentis de cada indicador dentro do setor (pesos iguais); maior é melhor',
                            'Aquisição de imobilizado e intangível dividida pela depreciação e amortização do ano (DFP 2025): acima de 1 = renovando/expandindo a base de ativos',
                            'CVM: ITR e DFP entregues até 22/10/2026 (últimos 12 meses); mesmas contas da plataforma',
                            'Varejo e construção civil (parametros.yaml); saúde e educação e transporte e logística retirados desta lista em 22/10/2026',
                            'Itaúsa (ITSA4) é controladora do Itaú (ITUB4): os dois juntos são a mesma exposição']}).to_excel(xw, sheet_name='Metodologia', index=False)
    for ws in xw.book.worksheets:
        ws.freeze_panes = 'A2'
        for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='0F181F')
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = max(10, min(60, max(len(str(c.value or '')) for c in col) + 2))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float): c.number_format = '0.0%' if str(ws.cell(1, c.column).value) in PCT else '0.00'
for s, g in T.groupby('Setor', sort=False): print(s, list(zip(g.pos, g.ticker)))
